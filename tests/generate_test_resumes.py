"""
Generate two dummy test resumes:
  - test_resume_A.docx  → title/date on SAME line (clean layout)
  - test_resume_B.pdf   → title/date on SEPARATE lines (messy layout)
    + uses non-standard section header names to stress-test alias map

Run directly: python tests/generate_test_resumes.py
"""

import pathlib
import sys

# ── helpers ─────────────────────────────────────────────────────────────────

TESTS_DIR = pathlib.Path(__file__).parent
RESUMES_DIR = TESTS_DIR / "sample_resumes"
RESUMES_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================================
# RESUME A — DOCX — title and date on the SAME line
# ============================================================================

RESUME_A_DATA = {
    "header": [
        "Alexandra Rivera",
        "alex.rivera@email.com | +1 (415) 555-0192 | linkedin.com/in/alexrivera",
        "San Francisco, CA",
    ],
    "sections": {
        "Technical Skills": [
            "Python, FastAPI, React, Node.js, PostgreSQL, Docker, AWS, Git",
            "TensorFlow, PyTorch, scikit-learn, Pandas, NumPy",
        ],
        "Experience": [
            "Senior Software Engineer, Stripe                    2021 - 2024",
            "• Led migration of payment processing service to microservices architecture.",
            "• Reduced API latency by 40% using Redis caching and async I/O.",
            "• Mentored a team of 4 junior engineers.",
            "",
            "Software Engineer, Lyft                            2019 - 2021",
            "• Developed real-time ride-matching algorithms using Python and Kafka.",
            "• Improved driver ETA accuracy by 15% via ML model integration.",
        ],
        "Education": [
            "B.S. Computer Science, Stanford University          2015 - 2019",
            "GPA 3.8 / 4.0, Honors",
        ],
        "Projects": [
            "ResumeAI",
            "• Built an end-to-end resume parsing pipeline using spaCy and FastAPI.",
            "• Achieved 92% field-extraction accuracy on a dataset of 1,000 resumes.",
            "",
            "TradeBot",
            "• Developed an algorithmic trading bot with Python and Alpaca API.",
        ],
        "Certifications": [
            "AWS Certified Solutions Architect – Associate (2023)",
            "Google Professional Data Engineer (2022)",
        ],
    },
}


def build_docx(data: dict, out_path: pathlib.Path):
    from docx import Document

    doc = Document()

    # Header block
    for line in data["header"]:
        doc.add_paragraph(line)

    doc.add_paragraph("")

    for section_title, lines in data["sections"].items():
        doc.add_heading(section_title, level=1)
        for line in lines:
            doc.add_paragraph(line)
        doc.add_paragraph("")

    doc.save(str(out_path))
    print(f"[OK] Created {out_path}")


# ============================================================================
# RESUME B — PDF — title and date on SEPARATE lines + non-standard headers
# ============================================================================

RESUME_B_LINES = """\
Jordan Kim
jordan.kim@protonmail.com
+1 (312) 555-0847
linkedin.com/in/jordan-kim-dev

PROFESSIONAL EXPERIENCE

Machine Learning Engineer
Uber AI Labs
2022 - 2024
• Designed and deployed a real-time fraud detection model serving 10M+ requests/day.
• Reduced false-positive rate by 22% through feature engineering and SHAP analysis.
• Collaborated with data platform team to migrate training pipelines to Spark.

Junior Data Scientist
Accenture Federal Services
2020 - 2022
• Built NLP pipelines to classify government contracts using scikit-learn and NLTK.
• Delivered weekly dashboards in Tableau for 3 federal agency clients.

CORE COMPETENCIES

Python, R, SQL, TensorFlow, PyTorch, scikit-learn, XGBoost, LightGBM
Docker, Kubernetes, AWS, GCP, Kafka, Spark
Git, Linux, REST

ACADEMIC BACKGROUND

Ph.D. (in progress) Computational Linguistics, University of Chicago
2022 - present

M.S. Data Science, Northwestern University
2018 - 2020

B.A. Mathematics, Carleton College
2014 - 2018

NOTABLE PROJECTS

SentimentStream
• Real-time Twitter sentiment analysis pipeline using Kafka, Spark Streaming, and BERT.
• Deployed on GCP with auto-scaling Kubernetes pods.

PolicyBot
• RAG-based chatbot for querying US federal policy documents using LangChain and OpenAI.

PROFESSIONAL CERTIFICATIONS

Google Professional Machine Learning Engineer (2023)
AWS Certified Machine Learning – Specialty (2021)
"""


def build_pdf(text: str, out_path: pathlib.Path):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import inch

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=LETTER,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )
    styles = getSampleStyleSheet()
    story = []

    for line in text.splitlines():
        if line.strip() == "":
            story.append(Spacer(1, 0.1 * inch))
        else:
            # Escape XML-special characters for ReportLab
            safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe, styles["Normal"]))

    doc.build(story)
    print(f"[OK] Created {out_path}")


# ============================================================================
# Main
# ============================================================================

if __name__ == "__main__":
    docx_path = RESUMES_DIR / "test_resume_A.docx"
    pdf_path = RESUMES_DIR / "test_resume_B.pdf"

    print("Generating test resumes …")
    build_docx(RESUME_A_DATA, docx_path)
    build_pdf(RESUME_B_LINES, pdf_path)
    print("Done. Resumes saved to:", RESUMES_DIR)
